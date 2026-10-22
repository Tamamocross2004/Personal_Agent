import csv
import io
import json

import pytest

from folder_assistant.exports import export_job,safe_csv
from test_folder_index import library,sync
from test_folder_answers import service,ask,finished


def test_export_roundtrip_formula_safety_and_unchanged_history(library):
    docs,folders,fid,index,embedding,_=library
    (docs/'a.txt').write_bytes(b'=danger alpha')  # Literal source value tests formula safety with a valid quote anchor.
    sync(library);jobs,chat=service(library)
    value=finished(jobs,ask(library,jobs));jid=value['id'];calls=len(embedding.calls)
    result=export_job(jobs,jid);output=folders.root/'outputs'/jid
    assert result['directory']==str(output) and not result['cached']
    csv_rows=list(csv.reader(io.StringIO((output/'comparison.csv').read_text('utf-8-sig'))))
    assert csv_rows[1][2]=="'=danger alpha" and 'SHA256' in csv_rows[1][4]
    from openpyxl import load_workbook
    book=load_workbook(output/'comparison.xlsx',data_only=False)
    assert book['comparison']['C2'].value=='=danger alpha' and book['comparison']['C2'].data_type=='s'
    book.close()
    source=json.loads((output/'sources.json').read_text('utf-8'))
    assert source['generation']==value['inputs']['generation'] and source['scope_files'][0]['file_sha256']
    assert source['delivered_evidence']==value['evidence']
    assert '未正式审稿' in (output/'summary.md').read_text('utf-8')
    assert json.loads((output/'record.json').read_text('utf-8'))==value
    folders.revoke(fid)
    assert export_job(jobs,jid)['cached'] and len(chat.calls)==1 and len(embedding.calls)==calls
    (output/'summary.md').write_text('user changed export',encoding='utf-8')
    with pytest.raises(ValueError,match='发生变化'):export_job(jobs,jid)
    assert (output/'summary.md').read_text('utf-8')=='user changed export'


@pytest.mark.parametrize('text',['=cmd()','+SUM(A1)','-1+2','@x','\t=1',' \r=1','\ufeff=1'])
def test_csv_dangerous_prefixes(text):assert safe_csv(text)=="'"+text


def test_blank_unknown_and_zero_are_distinct_in_exports(library):
    _,folders,_,_,_,_=library
    sync(library);jobs,_=service(library);job=finished(jobs,ask(library,jobs));jid=job['id']
    rows=[{'subject':'a','field':'quantity','value':'0','status':'stated','citations':['C1']},
          {'subject':'b','field':'quantity','value':None,'status':'unknown','citations':[]},
          {'subject':'c','field':'quantity','value':'','status':'blank','citations':['C1']}]
    job['result']['answer']['comparison']=rows;jobs.save(job)
    export_job(jobs,jid)
    from openpyxl import load_workbook
    book=load_workbook(folders.root/'outputs'/jid/'comparison.xlsx',data_only=False)
    sheet=book['comparison']
    assert sheet['C2'].value=='0' and sheet['C3'].value is None and sheet['C4'].value is None
    assert [sheet[f'D{i}'].value for i in (2,3,4)]==['stated','unknown','blank']
    book.close()


def test_table_rows_preserve_context_and_real_cells(library):
    docs,_,fid,index,_,_=library
    from openpyxl import Workbook
    book=Workbook();sheet=book.active;sheet.title='材料';sheet.append(['项目','数量']);sheet.append(['a',0]);sheet.append(['b','=1+2']);book.save(docs/'table.xlsx');book.close()
    sync(library)
    result=index.search(fid,'数量',mode='keyword')
    row=next(h for h in result['hits'] if h['source_location'].get('row')==2)
    assert row['source_location']['cells']==['A2','B2'] and '"value": 0' in row['text']
    source=index.citation(fid,result['generation'],row['doc_id'],row['source_block_id'])
    assert source['block']['source_blocks'][1]['location']['cell']=='B2'
    assert source['block']['source_blocks'][1]['value']==0


def test_unreadable_manifest_blocks_sync_and_keeps_previous(library):
    _,folders,fid,index,embedding,_=library
    old=sync(library);scan=folders.scan(fid);scan['entries'][0]['status']='unreadable'
    with folders.connect() as db:db.execute('UPDATE scans SET payload=? WHERE id=?',(json.dumps(scan),scan['id']))
    calls=len(embedding.calls)
    with pytest.raises(ValueError,match='不可读取'):index.sync(fid,scan['id'])
    assert index.store(fid).generation()['id']==old and len(embedding.calls)==calls
