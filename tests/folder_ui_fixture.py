"""Generate synthetic parser fixtures, never read user documents."""
from pathlib import Path
import sys
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject

root=Path(sys.argv[1]);root.mkdir(parents=True,exist_ok=True)
d=Document();d.add_paragraph('合同样例：交付前须核对清单。');t=d.add_table(rows=2,cols=2)
t.cell(0,0).text='字段';t.cell(0,1).text='记录';t.cell(1,0).text='状态';t.cell(1,1).text='待确认'
d.save(root/'说明.docx')
w=Workbook();s=w.active;s.title='台账';s.append(['项目','数量','公式']);s.append(['甲',0,'=1+2']);s.merge_cells('A3:B3');s['A3']='合并记录';w.save(root/'台账.xlsx')
(root/'数据.csv').write_text('名称,数值\n甲,0\n乙,\n',encoding='utf-8')
(root/'分页.txt').write_text('\n'.join('第'+str(i)+'行资料' for i in range(1,64)),encoding='utf-8')
(root/'损坏.docx').write_bytes(b'corrupt-office-file')
writer=PdfWriter();page=writer.add_blank_page(width=200,height=200)
font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
content=DecodedStreamObject();content.set_data(b'BT /F1 12 Tf 10 100 Td (Local PDF source) Tj ET')
page[NameObject('/Contents')]=writer._add_object(content);writer.write(root/'说明.pdf')
