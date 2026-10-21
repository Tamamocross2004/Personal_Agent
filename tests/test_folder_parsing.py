import hashlib,io,json,zipfile
from pathlib import Path
import pytest
from folder_assistant.parsing import extract
from folder_assistant.store import FolderStore

def test_text_locations_and_csv_multiline_records():
    raw='甲\r\n乙\n'.encode()
    result=extract(raw,'.txt')
    assert result['blocks'][1]['location']['start']==3
    assert result['blocks'][1]['text']=='乙\n'
    csv=extract('名称,数值,说明\r\n甲,0,"两行\n内容"\r\n乙,,\r\n'.encode(),'.csv')['blocks']
    assert csv[1]['location']['physical_line_start']==2 and csv[1]['location']['physical_line_end']==3
    assert csv[1]['cells'][1]['value']=='0' and csv[2]['cells'][1]['value']==''

def test_docx_source_order_tables_and_merged_cells():
    from docx import Document
    d=Document();d.add_paragraph('先读此段');t=d.add_table(rows=1,cols=2);t.cell(0,0).merge(t.cell(0,1)).text='合并内容';d.add_paragraph('最后一段')
    out=io.BytesIO();d.save(out);r=extract(out.getvalue(),'.docx')
    assert [b['location']['kind'] for b in r['blocks']]==['paragraph','table_cell','table_cell','paragraph']
    assert r['blocks'][2]['merged_origin']=={'row':1,'column':1}
    assert all('page' not in b['location'] for b in r['blocks'])

def workbook_bytes():
    from openpyxl import Workbook
    w=Workbook();s=w.active;s.title='金额';s.append(['项目','金额','计算']);s.append(['甲',0,'=1+2']);s.append(['乙',None,None]);s.merge_cells('A4:B4');s['A4']='合并';s.row_dimensions[3].hidden=True;s.column_dimensions['B'].hidden=True
    h=w.create_sheet('隐藏');h.sheet_state='hidden';h['A1']='隐藏记录'
    stream=io.BytesIO();w.save(stream);return stream.getvalue()

def test_xlsx_zero_blank_formula_cache_merges_and_hidden():
    raw=workbook_bytes();r=extract(raw,'.xlsx');cells={(b['location']['sheet'],b['location']['cell']):b for b in r['blocks']}
    assert cells['金额','B2']['value']==0 and cells['金额','B3']['value'] is None
    assert cells['金额','C2']['type']=='formula_cached_missing' and cells['金额','C2']['formula']=='=1+2'
    assert cells['金额','B4']['merged_range']=='A4:B4'
    assert cells['金额','B3']['hidden_row'] and cells['金额','B2']['hidden_column']
    assert cells['隐藏','A1']['sheet_state']=='hidden'
    source=zipfile.ZipFile(io.BytesIO(raw));target=io.BytesIO()
    with zipfile.ZipFile(target,'w') as archive:
        for name in source.namelist():
            data=source.read(name)
            if name=='xl/worksheets/sheet1.xml':data=data.replace(b'<f>1+2</f><v></v>',b'<f>1+2</f><v>3</v>')
            archive.writestr(name,data)
    result=extract(target.getvalue(),'.xlsx')
    cell=next(b for b in result['blocks'] if b['location']=={'kind':'spreadsheet_cell','sheet':'金额','cell':'C2'})
    assert cell['value']==3 and cell['type']=='formula_cached'

def test_pdf_page_and_blank_scan_state():
    from pypdf import PdfWriter
    writer=PdfWriter();writer.add_blank_page(width=100,height=100);buffer=io.BytesIO();writer.write(buffer)
    with pytest.raises(ValueError,match='文字'):extract(buffer.getvalue(),'.pdf')

@pytest.mark.parametrize('suffix',['.docx','.xlsx','.pdf'])
def test_corrupt_document_is_not_treated_as_empty_success(suffix):
    with pytest.raises(Exception):extract(b'not a document',suffix)

def test_zip_entity_input_rejected():
    s=io.BytesIO()
    with zipfile.ZipFile(s,'w') as z:z.writestr('word/document.xml','<!DOCTYPE a [<!ENTITY x "hidden">]><a>&x;</a>')
    with pytest.raises(Exception):extract(s.getvalue(),'.docx')

def test_preview_subprocess_cache_version_pagination_and_revoke(tmp_path):
    source=tmp_path/'source';source.mkdir();file=source/'data.csv';file.write_text('name,value\nalpha,0\nbeta,\n',encoding='utf-8')
    before=file.read_bytes();store=FolderStore(tmp_path/'state');key=store.grant(source)['id'];scan=store.scan(key);fid=scan['entries'][0]['id']
    first=store.preview(key,scan['id'],fid,limit=1)
    assert first['status']=='ready' and not first['cache_hit'] and first['total_blocks']==3
    second=store.preview(key,scan['id'],fid,offset=1,limit=1)
    assert second['cache_hit'] and second['blocks'][0]['cells'][1]['value']=='0'
    assert second['file_sha256']==hashlib.sha256(before).hexdigest() and file.read_bytes()==before
    store.revoke(key)
    with pytest.raises(ValueError,match='授权'):store.preview(key,scan['id'],fid)

def test_bad_encoding_and_changed_file_states(tmp_path):
    s=tmp_path/'source';s.mkdir();f=s/'bad.csv';f.write_bytes(b'\xff\xff')
    store=FolderStore(tmp_path/'state');key=store.grant(s)['id'];scan=store.scan(key);fid=scan['entries'][0]['id']
    assert store.preview(key,scan['id'],fid)['status']=='unsupported_encoding'
    f.write_text('new contents')
    with pytest.raises(ValueError,match='变化'):store.preview(key,scan['id'],fid)
