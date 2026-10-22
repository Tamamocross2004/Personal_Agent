"""Write-once project edition exports; original answers are grouped, never synthesized."""
import io,json,re
from openpyxl import Workbook
from knowledge.documents import sha
from knowledge.locking import lease
from .store import local_path
from .exports import NAMES,csv_bytes,md


def export_version(projects,vid):
    if not re.fullmatch('[a-f0-9]{32}',vid):raise ValueError('结果版本ID无效')
    version=projects._read('folder_project_versions',vid)
    signature=sha(json.dumps(version,sort_keys=True,ensure_ascii=False))
    root=local_path(projects.folders.root);output=root/'project-outputs'/vid
    for p in (root/'project-outputs',output):
        if p.exists() or p.is_symlink():
            if local_path(p)!=p or not p.is_dir():raise ValueError('项目导出路径无效')
        else:p.mkdir()
    with lease(output/'.export.lock'):
        manifest=output/'export.json'
        if manifest.is_symlink():raise ValueError('项目导出清单不得为链接')
        if manifest.exists():
            saved=json.loads(manifest.read_text('utf-8'))
            if saved['version_sha256']!=signature:raise ValueError('项目版本变化，保留旧导出')
            for name,h in saved['files'].items():
                if name not in NAMES or local_path(output/name)!=output/name or sha((output/name).read_bytes())!=h:
                    raise ValueError('项目导出文件变化，未覆盖')
            return {'directory':str(output),'files':list(saved['files']),'cached':True}
        if any((output/n).exists() or (output/n).is_symlink() for n in NAMES):raise ValueError('项目导出不完整，保留原文件，请另存核对')
        groups=[projects.group(e) for e in version['entries']]
        rows=[['subject','field','value','status','citations_and_source_versions','subject_type','subject_quote','value_quote','job_id','batch_number']]
        sources=[['job_id','batch_number','citation_id','relative_path','generation','file_sha256','source_location','role']]
        lines=['# '+md(version['config']['title']),'',md(version['config']['objective']),
               '版本：'+vid+'；父版本：'+str(version['parent_id']),
               '状态：'+version['status']+'；未正式审稿；语义未认证。',md(version['meaning']),
               '字段：'+md('、'.join(version['config']['fields'])),
               '覆盖：'+md(json.dumps(version['coverage'],ensure_ascii=False)),
               '变化：'+md(json.dumps(version['changes'],ensure_ascii=False)),
               '以下仅汇集未截断的可用批次，按原来源保留；没有业务对象去重、跨文件合并或语义认证。','']
        for g in groups:
            evidence={e['citation_id']:e for e in g['evidence']}
            lines.extend(['## '+g['job_id']+' 第'+str(g['batch'])+'批','固定索引：'+g['generation']])
            lines.append('本批要求：'+md(g['question']))
            for f in g['answer']['findings']:
                lines.append('- '+md(f['kind']+'：'+f['statement']))
                for label,key in [('假设','assumptions'),('限制','limitations')]:
                    for text in f.get(key,[]):lines.append('  - '+label+'：'+md(text))
                if f.get('relation'):lines.append('  - 比较条件：'+md(json.dumps(f['relation'],ensure_ascii=False)))
                for q in f.get('support',[]):lines.append('  - '+md(q['citation_id']+' '+q['quote']))
            for text in g['answer']['unknowns']:lines.append('- 待核对：'+md(text))
            for r in g['answer']['comparison']:
                refs='; '.join(f"{g['job_id']}/{cid} {evidence[cid]['name']} SHA256 {evidence[cid]['file_sha256']} 索引 {g['generation']} {json.dumps(evidence[cid]['source_location'],ensure_ascii=False)}" for cid in r['citations'])
                rows.append([r['subject'],r['field'],r['value'],r['status'],refs,r['subject_type'],
                             json.dumps(r['subject_quote'],ensure_ascii=False),json.dumps(r['value_quote'],ensure_ascii=False),g['job_id'],g['batch']])
            for e in g['evidence']:
                sources.append([g['job_id'],g['batch'],e['citation_id'],e['name'],g['generation'],e['file_sha256'],json.dumps(e['source_location'],ensure_ascii=False),e.get('role')])
        book=Workbook();sheet=book.active;sheet.title='comparison'
        for name,data in [('comparison',rows),('sources',sources)]:
            s=sheet if name=='comparison' else book.create_sheet(name);s.freeze_panes='A2'
            for row in data:
                s.append(row)
                for c in s[s.max_row]:
                    if isinstance(c.value,str):c.data_type='s'
        buffer=io.BytesIO();book.save(buffer);book.close()
        payload={'summary.md':'\n'.join(lines).encode('utf-8'),'comparison.csv':csv_bytes(rows),
                 'comparison.xlsx':buffer.getvalue(),'sources.csv':csv_bytes(sources),
                 'sources.json':json.dumps({'version':version,'groups':[{k:v for k,v in g.items() if k!='answer'} for g in groups]},ensure_ascii=False,indent=2).encode('utf-8'),
                 'record.json':json.dumps({'version':version,'groups':groups},ensure_ascii=False,indent=2).encode('utf-8')}
        for name,data in payload.items():
            with (output/name).open('xb') as f:f.write(data)
        with manifest.open('x',encoding='utf-8') as f:json.dump({'version_sha256':signature,'files':{n:sha(b) for n,b in payload.items()}},f,indent=2)
        return {'directory':str(output),'files':list(payload),'cached':False}
