"""Read-only task exports. User text is never emitted as executable spreadsheet formulas."""
import csv
import io
import json
from pathlib import Path
import re
from uuid import uuid4

from knowledge.documents import sha
from knowledge.locking import lease
from .store import local_path


NAMES={'summary.md','comparison.csv','comparison.xlsx','sources.json','sources.csv','record.json'}


def safe_csv(value):
    value='' if value is None else str(value)
    return "'"+value if value.lstrip(' \t\r\n\ufeff').startswith(('=','+','-','@')) else value


def csv_bytes(rows):
    stream=io.StringIO(newline='')
    writer=csv.writer(stream,quoting=csv.QUOTE_ALL)
    writer.writerows([[safe_csv(v) for v in row] for row in rows])
    return stream.getvalue().encode('utf-8-sig')


def md(value):
    text=str(value).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')
    return re.sub(r'([\\`*_{}\[\]()#!|])',r'\\\1',text)


def export_job(jobs,jid):
    job=jobs.get(jid)
    if job['kind']!='answer' or job['status'] in ('running','queued'): raise ValueError('请等待短答结束后导出；失败或中断记录也可导出核对')
    if not re.fullmatch('[a-f0-9]{32}',jid): raise ValueError('任务ID无效')
    root=local_path(jobs.folders.root)
    output=root/'outputs'/jid
    # Reject any pre-existing link before creating or writing into the fixed task directory.
    for p in (root/'outputs',output):
        if p.exists() or p.is_symlink():
            if local_path(p)!=p or not p.is_dir(): raise ValueError('导出路径无效')
        else:p.mkdir()
    if not local_path(output).is_relative_to(root): raise ValueError('导出越出任务目录')
    with lease(output/'.export.lock'):
        manifest=output/'export.json'
        if manifest.is_symlink(): raise ValueError('导出清单不得为链接')
        signature=sha(json.dumps(job,sort_keys=True,ensure_ascii=False))
        if manifest.exists():
            saved=json.loads(manifest.read_text('utf-8'))
            if saved['record_sha256']!=signature: raise ValueError('任务记录发生变化；原导出已保留，未覆盖')
            for name,digest in saved['files'].items():
                if name not in NAMES or local_path(output/name)!=(output/name) or sha((output/name).read_bytes())!=digest:
                    raise ValueError('已有导出文件发生变化；保留原文件，未自动覆盖')
            return {'directory':str(output),'files':list(saved['files']),'cached':True}
        answer=(job.get('result') or {}).get('answer') or {}
        batches=job.get('batches',[])
        if batches:
            # Concatenate for export only, with batch labels; no synthesized merged answer.
            answer={'schema_version':3,'findings':[],'comparison':[],'unknowns':[]}
            for batch in batches:
                body=(batch.get('result') or {}).get('answer') or {}
                for key in ('findings','comparison'):
                    answer[key].extend({**r,'export_batch':batch['number']} for r in body.get(key,[]))
                answer['unknowns'].extend(f"第{batch['number']}批：{s}" for s in body.get('unknowns',[]))
        evidence={e['citation_id']:e for e in job.get('evidence',[])}
        source={'task_id':jid,'generation':job['inputs']['generation'],'review_status':'not_reviewed',
                'scope_files':job['inputs']['source_snapshot'],'delivered_evidence':list(evidence.values()),
                'retrieval':job.get('retrieval'),
                'meaning':'范围文件不等于模型阅读全文；引用/定位仅认证保存来源，不认证结论正确。'}
        from .scope import scope_view
        source['scope_view']=scope_view(jobs,jid,compact=True)
        if 'coverage' in job:source['coverage']=job['coverage']
        if 'batch_decisions' in job:
            source['batch_decisions']=job['batch_decisions']
            source['output_policy']={k:job['inputs'].get(k) for k in ('output_tokens','effective_output_cap','batch_mode','batch_size','expected_fields','retry_of','range_end')}
        if 'strength' in job['inputs']:
            source.update(strength=job['inputs']['strength'],model=job['inputs'].get('model'))
        def refs(ids):
            return '; '.join(cid+' '+evidence[cid]['name']+' '+json.dumps(evidence[cid]['source_location'],ensure_ascii=False)+
                             ' SHA256 '+evidence[cid]['file_sha256'] for cid in ids)
        comparison=[['subject','field','value','status','citations_and_source_versions']]
        v2=answer.get('schema_version',1)>=2
        if v2:comparison[0].extend(['subject_type','subject_quote','value_quote'])
        if batches:comparison[0].append('batch_number')
        for row in answer.get('comparison',[]):
            values=[row['subject'],row['field'],row['value'],row['status'],refs(row['citations'])]
            if v2:values.extend([row.get('subject_type','legacy'),json.dumps(row.get('subject_quote'),ensure_ascii=False),json.dumps(row.get('value_quote'),ensure_ascii=False)])
            if batches:values.append(row['export_batch'])
            comparison.append(values)
        sources=[['relative_path','file_sha256','scan_id','generation','delivered_citations','meaning']]
        for doc in source['scope_files']:
            ids=[cid for cid,e in evidence.items() if e['doc_id']==doc['id']]
            sources.append([doc['name'],doc['file_sha256'],doc['scan_id'],source['generation'],','.join(ids),
                            '实际送达有限片段' if ids else '仅在指定范围，不代表已阅读'])
        lines=['# 文件夹资料短答','',md(job['inputs']['query']),'',
               '未正式审稿；执行状态：'+job['status']+'；内容状态：'+(job.get('result') or {}).get('content_status','未形成答案'),
               '固定索引：'+source['generation'],'']
        if job.get('message'): lines.extend([md(job['message']),''])
        view=source['scope_view']
        lines.extend(['## 范围与任务链（只读记录快照）',md(view['meaning']),''])
        if view['mode']=='sequential':
            lines.extend(['按片段去重状态：'+md(json.dumps(view['counts'],ensure_ascii=False)),
                          md(view['counts_meaning'])])
            for span in view['runs']:
                lines.append(md(f"[{span['start']}, {span['end']}) {span['name']}：{span['label']}"+('；保留此前失败尝试' if span['failed_before'] else '')))
            for item in view['chain']:
                lines.append(md(f"任务 {item['task_id']} · {item['operation']} · 已知Token {item['known_tokens']} · 未知用量事件 {item['unknown_usage_events']}"))
        else:lines.append('局部检索，不能计算全文阅读率；确认返回的有限证据：'+str(view['confirmed_evidence'])+'段。')
        for warning in view['warnings']:lines.append('范围记录不完整：'+md(warning))
        lines.append('')
        if 'strength' in job['inputs']:lines.extend(['模型强度：'+md(job['inputs']['strength'])+'；实际effort：'+md((job['inputs'].get('model') or {}).get('reasoning_effort') or '默认'),''])
        if 'batch_decisions' in job:
            lines.extend(['输出与分批设置：'+md(json.dumps(source['output_policy'],ensure_ascii=False)),
                          '分批只作资源估算，不认证完整性或语义正确；原失败和费用保留，拆小重做属于新的付费请求。',''])
        if 'coverage' in job:
            lines.extend(['顺序分批整理：以下结果按批次保留，未做跨批合并；续读的前次结果保留在原任务。',
                          '提取范围送达不认证完整抽取或语义正确。',md(json.dumps(job['coverage'],ensure_ascii=False)),''])
            for batch in batches:
                if not (batch.get('result') or {}).get('answer'):
                    lines.append(f"第{batch['number']}批未形成可用结构化结果；原始响应见record.json。")
        diagnostic_rows=[(f"第{b['number']}批",b.get('output_diagnostics')) for b in batches] if batches else [('',job.get('output_diagnostics'))]
        for label,diagnostic in diagnostic_rows:
            if diagnostic:
                lines.extend([md(label+' 输出诊断：'+diagnostic['message']),md(diagnostic['suggestion']),
                              '用量状态：'+md(diagnostic['usage_status'])+'；不认证语义正确性。',''])
        for finding in answer.get('findings',[]):
            if batches:lines.append('第'+str(finding['export_batch'])+'批：')
            lines.extend(['- '+md(finding['kind']+'：'+finding['statement'])+' '+md(' '.join('['+c+']' for c in finding['citations'])),''])
            if v2:
                for label,key in [('假设（未证实）','assumptions'),('判断限制','limitations')]:
                    for value in finding.get(key,[]):lines.append('  - '+label+'：'+md(value))
                if finding.get('relation'):
                    r=finding['relation'];lines.append('  - 比较条件：'+md(json.dumps(r,ensure_ascii=False)))
                for q in finding.get('support',[]):lines.append('  - 依据原句 '+md(q['citation_id']+'：'+q['quote']))
        if v2:
            lines.extend(['','结构/原句匹配不认证事实或矛盾判断。'])
            for warning in (job.get('result') or {}).get('contract_checks',{}).get('warnings',[]):
                lines.append('- 待核对：'+md(warning['message']))
        for unknown in answer.get('unknowns',[]):lines.append('- 待核对：'+md(unknown))
        lines.extend(['','## 实际送达证据',''])
        for cid,e in evidence.items():lines.extend([md('['+cid+'] '+e['name']+' '+json.dumps(e['source_location'],ensure_ascii=False)),
             'SHA256 '+e['file_sha256'],'',*['> '+md(line) for line in e['text'].splitlines()],''])
        payload={'summary.md':'\n'.join(lines).encode('utf-8'), 'comparison.csv':csv_bytes(comparison),'sources.csv':csv_bytes(sources),
                 'sources.json':json.dumps(source,ensure_ascii=False,indent=2).encode('utf-8'),
                 'record.json':json.dumps(job,ensure_ascii=False,indent=2).encode('utf-8')}
        if any((output/name).exists() or (output/name).is_symlink() for name in NAMES):
            raise ValueError('该任务有不完整导出；已保留文件，请核对后手动另存，未覆盖')
        from openpyxl import Workbook
        workbook=Workbook();sheet=workbook.active;sheet.title='comparison'
        for row in comparison:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        sheet.freeze_panes='A2'
        info=workbook.create_sheet('sources')
        for row in sources:
            info.append(row)
            for cell in info[info.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        buffer=io.BytesIO();workbook.save(buffer);workbook.close();payload['comparison.xlsx']=buffer.getvalue()
        # Write-once files. An interrupted partial export is retained for diagnosis, never replaced.
        for name,content in payload.items():
            with (output/name).open('xb') as stream: stream.write(content)
        with manifest.open('x',encoding='utf-8') as stream:
            json.dump({'record_sha256':signature,'files':{name:sha(data) for name,data in payload.items()}},stream,ensure_ascii=False,indent=2)
        return {'directory':str(output),'files':list(payload),'cached':False}
