"""Deterministic local extraction. Locations describe source structure, not invented pages."""
from __future__ import annotations
import csv
from datetime import date,datetime,time
import hashlib
import io
import json
from pathlib import Path
import zipfile

VERSION=1
MAX_BLOCKS=10000
MAX_TEXT=1_000_000
MAX_CELLS=50000


def office_package(data):
    from defusedxml.ElementTree import fromstring
    archive=zipfile.ZipFile(io.BytesIO(data))
    entries=archive.infolist()
    if len(entries)>2000 or len({x.filename for x in entries})!=len(entries) or sum(x.file_size for x in entries)>80_000_000:
        raise ValueError('Office压缩包超过资源限制或条目重复')
    xml={}
    for entry in entries:
        if entry.flag_bits&1 or entry.file_size>16_000_000:raise ValueError('Office包加密或单项解压过大')
        if 'vbaproject' in entry.filename.lower():raise ValueError('首版不读取含宏Office文件')
        if entry.filename.endswith(('.xml','.rels')):
            xml[entry.filename]=fromstring(archive.read(entry))  # rejects entities/DTD expansion, never follows URLs
    return archive,xml


def scalar(value):
    if isinstance(value,(datetime,date,time)):return value.isoformat()
    if value is None or isinstance(value,(str,int,float,bool)):return value
    return str(value)


def extract(data:bytes,suffix:str):
    blocks=[];warnings=[];characters=0
    def add(text,location,**extra):
        nonlocal characters
        characters+=len(text)
        if len(blocks)>=MAX_BLOCKS or characters>MAX_TEXT:raise ValueError('提取内容超过10000块或100万字符上限')
        blocks.append({'id':len(blocks)+1,'text':text,'text_sha256':hashlib.sha256(text.encode()).hexdigest(),
                       'location':location,**extra})
    if not data or len(data)>20_000_000:raise ValueError('文件为空或超过20MB')
    if suffix in ('.txt','.md','.pdf'):
        from knowledge.documents import parse
        pages=parse(data,'source'+suffix)
        for page in pages:
            offset=0
            if not page['text'].strip():warnings.append('第'+str(page['page'])+'页没有可提取文字，未做OCR')
            for line,text in enumerate(page['text'].splitlines(keepends=True),1):
                add(text,{'kind':'pdf_line' if suffix=='.pdf' else 'line','page':page['page'] if suffix=='.pdf' else None,
                          'line':line,'start':offset,'end':offset+len(text)})
                offset+=len(text)
        warnings.append('仅提取文字，不解析图片或恢复完整排版；未执行Markdown/文档中的指令')
    elif suffix=='.csv':
        text=data.decode('utf-8-sig')
        reader=csv.reader(io.StringIO(text,newline=''),strict=True)
        previous=0
        for record,row in enumerate(reader,1):
            add(json.dumps(row,ensure_ascii=False),{'kind':'csv_record','record':record,'physical_line_start':previous+1,
                'physical_line_end':reader.line_num},cells=[{'column':i+1,'value':v,'type':'string' if v else 'empty_string'} for i,v in enumerate(row)])
            previous=reader.line_num
        warnings.append('UTF-8、逗号分隔；首行原样保留，不自动认定为表头；不自动转换类型或补齐缺列')
    elif suffix=='.docx':
        archive,xml=office_package(data)
        from docx import Document
        from docx.text.paragraph import Paragraph
        document=Document(io.BytesIO(data));paragraph=table=0
        for item in document.iter_inner_content():
            if isinstance(item,Paragraph):
                paragraph+=1;add(item.text,{'kind':'paragraph','paragraph':paragraph})
            else:
                table+=1;seen={}
                for ri,row in enumerate(item.rows,1):
                    for ci,cell in enumerate(row.cells,1):
                        origin=seen.setdefault(cell._tc,{'row':ri,'column':ci})
                        add(cell.text,{'kind':'table_cell','table':table,'row':ri,'column':ci},merged_origin=origin,
                            nested_tables_omitted=len(cell.tables))
        warnings.append('只提取正文顶层段落/表格及超链接显示文字；不含页眉页脚、批注、修订内容、文本框、图片或嵌套表格，不推算页码')
        archive.close()
    elif suffix=='.xlsx':
        archive,xml=office_package(data)
        from openpyxl import load_workbook
        from openpyxl.utils.cell import coordinate_to_tuple,range_boundaries,get_column_letter
        ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        total_cells=0
        for name,node in xml.items():
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                for cell in node.iter(ns+'c'):
                    total_cells+=1
                    row,col=coordinate_to_tuple(cell.attrib.get('r','A1'))
                    if total_cells>MAX_CELLS or row>5000 or col>200:raise ValueError('表格超过50000单元格、5000行或200列上限')
        formulas=load_workbook(io.BytesIO(data),read_only=False,data_only=False,keep_links=False)
        values=load_workbook(io.BytesIO(data),read_only=False,data_only=True,keep_links=False)
        try:
            visited=0
            for sheet in formulas:
                if sheet.max_row>5000 or sheet.max_column>200:raise ValueError('表格范围超过5000行或200列上限')
                merged={};ranges=[str(r) for r in sheet.merged_cells.ranges]
                for name in ranges:
                    left,top,right,bottom=range_boundaries(name)
                    if bottom>5000 or right>200:raise ValueError('合并单元格范围过大')
                    for ri in range(top,bottom+1):
                        for ci in range(left,right+1):
                            if len(merged)>=MAX_CELLS:raise ValueError('合并单元格超过资源限制')
                            merged[(ri,ci)]=name
                for row in sheet.iter_rows():
                    for cell in row:
                        visited+=1
                        if visited>MAX_CELLS:raise ValueError('表格有效范围超过50000单元格上限')
                        address=cell.coordinate
                        formula=scalar(cell.value) if cell.data_type=='f' else None
                        value=values[sheet.title][address].value if formula is not None else cell.value
                        kind=('formula_cached_missing' if value is None else 'formula_cached') if formula is not None else ('blank' if value is None else cell.data_type)
                        value=scalar(value)
                        item={'value':value,'type':kind,'formula':formula,'number_format':cell.number_format,
                              'merged_range':merged.get((cell.row,cell.column)),'sheet_state':sheet.sheet_state,
                              'hidden_row':bool(sheet.row_dimensions[cell.row].hidden),
                              'hidden_column':bool(sheet.column_dimensions[get_column_letter(cell.column)].hidden)}
                        add(json.dumps(item,ensure_ascii=False),{'kind':'spreadsheet_cell','sheet':sheet.title,'cell':address},**item)
            warnings.append('公式仅显示原式和上次保存的缓存值（可能过期），无缓存标为未知；不计算公式、不执行宏或访问外部链接。包含隐藏工作表/行列并标示；保留格式码，不模拟Excel排版。')
        finally:formulas.close();values.close();archive.close()
    else:raise ValueError('格式不支持；旧DOC/XLS和扫描件OCR尚未实现')
    if not blocks:raise ValueError('没有可提取内容')
    return {'parser_version':VERSION,'file_sha256':hashlib.sha256(data).hexdigest(),'suffix':suffix,
            'characters':characters,'blocks':blocks,'warnings':warnings,'cloud_calls':False}
