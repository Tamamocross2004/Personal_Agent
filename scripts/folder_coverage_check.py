"""Isolated large-corpus structural diagnostic. NO real embedding or language calls.

The lexical result and exhaustive reading plan are measured separately. Local text
coverage is not extraction accuracy, model delivery, semantic recall or model quality.
"""
import argparse
import csv
import json
from pathlib import Path
import re
import sys
import time

PROJECT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(PROJECT),str(PROJECT/'tests')]
from agent.config import Settings
from folder_assistant.store import FolderStore
from folder_assistant.index import FolderIndex
from folder_assistant.reading import ReadingPlan
from knowledge.store import KnowledgeStore
from knowledge.documents import sha
from test_folder_index import Embedding,Vectors


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args();root=args.output.resolve()
    root.mkdir(parents=True,exist_ok=False);docs=root/'documents';docs.mkdir()
    for n in range(10):
        with (docs/f'ledger-{n:02}.csv').open('w',encoding='utf-8',newline='') as stream:
            writer=csv.writer(stream);writer.writerow(['device','cycle','unit','note'])
            for row in range(200):writer.writerow([f'Unit-{n*200+row:04}',row%17+1,'day','含零值0和多语言条件，原样保留'])
    for n in range(4):
        (docs/f'notes-{n}.md').write_text('# 范围与适用条件\n'+''.join(f'章节段落{i}：只适用于记录中的实际条件，不补缺项。\n' for i in range(80)),encoding='utf-8')
    from openpyxl import Workbook
    workbook=Workbook();sheet=workbook.active;sheet.append(['object','value','unit'])
    for row in range(200):sheet.append([f'Sheet-{row:03}',0 if row%3==0 else row,'kPa'])
    workbook.save(docs/'pressure.xlsx');workbook.close()
    from docx import Document
    document=Document();document.add_paragraph('范围说明：保留真实表格行号。');table=document.add_table(rows=31,cols=2)
    for row in range(31):table.cell(row,0).text=f'条目{row}';table.cell(row,1).text=f'条件{row}'
    document.save(docs/'conditions.docx')
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
    writer=PdfWriter()
    for page in range(3):
        target=writer.add_blank_page(width=600,height=800)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        target[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        content=DecodedStreamObject();content.set_data(f'BT /F1 12 Tf 40 750 Td (Page {page+1}: no inferred table join.) Tj 0 -30 Td (Local condition {page+1}.) Tj ET'.encode())
        target[NameObject('/Contents')]=writer._add_object(content)
    writer.write(docs/'pages.pdf')
    source_hashes={p.name:sha(p.read_bytes()) for p in docs.iterdir()}
    prereg={'paid_calls':0,'fake_providers':True,'sources':source_hashes,'expected_csv_records':2010,
            'query':'cycle','case':'header-only lexical query versus local sequential plan',
            'production_capacity_changes':False,'semantic_quality_assessed':False}
    (root/'plan.json').write_text(json.dumps(prereg,ensure_ascii=False,indent=2),'utf-8')
    settings=Settings(_env_file=None,siliconflow_api_key='explicit-test-fake',embedding_dimensions=2,knowledge_build_token_budget=100000)
    folders=FolderStore(root/'state');fid=folders.grant(docs)['id'];embedding=Embedding();vectors=Vectors()
    index=FolderIndex(folders,settings,factory=lambda s:KnowledgeStore(s,vectors=vectors,embedding_client=embedding))
    began=time.perf_counter();scan=folders.scan(fid);built=index.sync(fid,scan['id']);build_seconds=time.perf_counter()-began
    generation=built['active']['id'];began=time.perf_counter();result=index.search(fid,'cycle',generation=generation,mode='keyword');lexical_seconds=time.perf_counter()-began
    began=time.perf_counter();plan=ReadingPlan(index,fid,generation);offset=0;selected=[];context_chars=0;batch_chars=[]
    while offset<len(plan.segments):
        batch=plan.batch(offset);selected.extend(h for h in batch['hits'] if h['role']=='sequential')
        context_chars+=sum(len(h['text']) for h in batch['hits'] if h['role']!='sequential')
        batch_chars.append(sum(len(h['text']) for h in batch['hits']));offset=batch['next_offset']
    read_seconds=time.perf_counter()-began
    def objects(hits):return sorted(set(re.findall(r'Unit-\d{4}','\n'.join(h['text'] for h in hits))))
    expected={f'Unit-{n:04}' for n in range(2000)}
    actual=set(objects(selected));lexical=set(objects(result['hits']))
    assert actual==expected and [h['segment'] for h in selected]==list(range(len(plan.segments)))
    assert max(batch_chars)<=5500
    assert source_hashes=={p.name:sha(p.read_bytes()) for p in docs.iterdir()}
    report={'preregistration':prereg,'generation':generation,'sources_unchanged':True,'file_count':len(plan.documents),
            'indexed_chunks':built['active']['chunks'],'reading_plan':plan.summary(),
            'build_seconds_with_fake_embedding':round(build_seconds,3),'keyword_seconds':round(lexical_seconds,4),
            'local_sequential_seconds':round(read_seconds,4),'local_batches':len(batch_chars),'max_batch_characters':max(batch_chars),
            'repeated_context_characters':context_chars,
            'keyword_csv_object_coverage':{'found':len(lexical & expected),'expected':len(expected),'ratio':len(lexical & expected)/len(expected)},
            'local_sequential_csv_object_coverage':{'found':len(actual),'expected':len(expected),'ratio':1},
            'paid_calls':0,'real_model_tokens':None,'real_model_cost':None,'semantic_quality':'not_assessed',
            'limitation':'局部关键词覆盖低是此header-only案例的观察，不是RAG总体召回率；顺序计划遍历100%不等于模型完整抽取、内容正确或真实大库延迟。'}
    (root/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8')
    print(json.dumps({k:v for k,v in report.items() if k not in ('preregistration','reading_plan')},ensure_ascii=False))


if __name__=='__main__':main()
